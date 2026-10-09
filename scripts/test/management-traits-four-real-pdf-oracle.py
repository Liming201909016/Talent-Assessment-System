"""Independent rational/Excel/actual server-PDF assertions; synthetic owned rows only."""
from decimal import Decimal, ROUND_HALF_UP
from fractions import Fraction
import hashlib
import json
from pathlib import Path
import re
import sys

import pymupdf as fitz
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
DIRECTORY = ROOT / 'scripts/test/results/MTRa05eb3af5e49'
SPECS = [
    ('self_confidence','self',[6,19,33,47,60,66,78,92,101,107,-120,-133],Fraction(115,2)),
    ('emotional_stability','self',[10,24,38,52,68,82,96,-111,-124,-137],Fraction(55)),
    ('self_discipline','self',[11,-25,39,53,-69,84,-97,-112,125,138],Fraction(225,4)),
    ('sociality','interpersonal',[1,14,-28,-42,-56,-72,-87,-102,-115,-128],Fraction(205,4)),
    ('leadership','interpersonal',[3,16,31,44,58,74,89,104,-117,-130],Fraction(105,2)),
    ('interpersonal_sensitivity','interpersonal',[5,18,32,46,50,61,65,76,83,91,-106,-119,-132],Fraction(50)),
    ('cooperation','interpersonal',[13,27,41,55,71,86,99,-114,127,-140],Fraction(115,2)),
    ('planning','task',[4,17,21,30,45,59,75,81,90,105,118,-131],Fraction(215,4)),
    ('responsibility','task',[7,20,34,48,63,77,-94,-108,-121,-134],Fraction(235,4)),
    ('decisiveness','task',[-12,26,35,40,54,62,70,85,98,113,126,139],Fraction(215,4)),
    ('proactiveness','development',[2,15,29,43,-57,73,88,103,-116,129],Fraction(55)),
    ('learning','development',[-8,22,36,49,-64,79,93,-100,109,122,-135],Fraction(215,4)),
    ('innovation','development',[-9,23,37,-51,67,-80,95,110,123,-136],Fraction(50)),
]

def rounded(value, digits):
    return (Decimal(value.numerator) / Decimal(value.denominator)).quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP)

def level(value):
    return next(label for threshold,label in [(90,'excellent'),(70,'good'),(30,'qualified'),(10,'weak'),(0,'insufficient')] if value >= threshold)

def compact(text):
    return re.sub(r'\s+', '', text)

def main():
    rows = json.loads((DIRECTORY/'owned-sql.json').read_text('utf-8'))
    workbook = ROOT/'Go-based Refactored System/configs/export-templates/management-traits-002-test-content-v1.xlsx'
    assert hashlib.sha256(workbook.read_bytes()).hexdigest() == 'b0498249ae057e3aa1b798922ed2d53a6ca304811943b3b41f4f64f9b024e84c'
    sheets = openpyxl.load_workbook(workbook, data_only=True).worksheets
    output = []
    for index,row in enumerate(rows):
        questions = {q['number']:q for q in row['questions']}
        assert sorted(questions) == list(range(1,141))
        dims = {d['key']:d for d in row['dimensions']}
        scores = {}
        for order,(key,module,terms,norm) in enumerate(SPECS,1):
            total = 0
            for term in terms:
                q = questions[abs(term)]
                expected = [3,1,5,1+abs(term)%5][index]
                raw = 6-expected if term < 0 else expected
                assert q['reverse'] == (term < 0) and q['raw'] == raw and q['final'] == expected
                assert next(o['raw'] for o in q['options'] if o['sourceOptionId'] == q['selected']) == raw
                total += expected
            value = Fraction(25*(total-len(terms)),len(terms))
            d = dims[key]
            assert (d['order'],d['sum'],d['count'],d['answered']) == (order,total,len(terms),len(terms))
            assert Decimal(str(d['score'])) == rounded(value,6) and Decimal(str(d['norm'])) == rounded(norm,6) and d['level'] == level(value)
            scores[key] = value
        overall = sum(scores.values(),Fraction())/13
        assert Decimal(str(row['run']['score'])) == rounded(overall,6)
        assert Decimal(str(row['run']['norm'])) == rounded(Fraction(705,13),6)
        assert row['run']['level'] == level(overall) and row['run']['status'] == 'completed'
        for mod in row['modules']:
            children = [scores[key] for key,module,_,_ in SPECS if module == mod['key']]
            assert mod['count'] == len(children) and Decimal(str(mod['score'])) == rounded(sum(children,Fraction())/len(children),6)
        receipt = row['receipt']
        assert receipt['runId'] == row['run']['id'] and receipt['type'] == 'manual'
        assert receipt['userTime'] == row['run']['userTime'] == row['userTime']
        assert receipt['submitted'] == row['run']['submitted'] and receipt['started'] == row['started'] and receipt['deadline'] == row['deadline']
        report = row['report']
        assert hashlib.sha256(report['dataRaw'].encode()).hexdigest() == report['dataSha']
        dto = json.loads(report['dataRaw'])
        assert dto['runId'] == row['run']['id'] and dto['testOnly'] and len(dto['dimensions']) == 13 and len(dto['modules']) == 4
        assert dto['overall']['score'] == str(rounded(overall,2))
        ranking = sorted(range(13),key=lambda n:(-scores[SPECS[n][0]],n))
        low = sorted(range(13),key=lambda n:(scores[SPECS[n][0]],n))
        assert [x['Key'] for x in dto['highest']] == [SPECS[n][0] for n in ranking[:3]]
        assert [x['Key'] for x in dto['lowest']] == [SPECS[n][0] for n in low[:3]]
        file = DIRECTORY/('00201' if index < 2 else '00202')
        file = file.with_name(file.name+('-tester.pdf' if index%2 else '-candidate.pdf'))
        raw = file.read_bytes()
        assert len(raw) == report['fileBytes'] and hashlib.sha256(raw).hexdigest() == report['fileSha'] and raw.startswith(b'%PDF-')
        pdf = fitz.open(file)
        text = compact(''.join(p.get_text(sort=False) for p in pdf))
        assert 'TEST' in text and '不可作为人才决策依据' in text and '\ufffd' not in text
        selected = []
        for n,d in enumerate(dto['dimensions']):
            grade = ['excellent','good','qualified','weak','insufficient'].index(dims[d['key']]['level'])
            assert d['diagnosis'] == sheets[1].cell(n+3,4+grade).value
            assert d['advice'] == sheets[2].cell(n+3,3+grade).value
            assert d['score'] == str(rounded(scores[d['key']],2))
            selected += [d['diagnosis'],d['advice']]
        for item in dto['highest']+dto['lowest']:
            n = [s[0] for s in SPECS].index(item['Key'])
            grade = ['excellent','good','qualified','weak','insufficient'].index(dims[item['Key']]['level'])
            assert item['Text'] == sheets[0].cell(n+3,3+grade).value
            selected.append(item['Text'])
        grade = ['excellent','good','qualified','weak','insufficient'].index(row['run']['level'])
        assert dto['overall']['diagnosis'] == sheets[3].cell(grade+2,3).value
        assert dto['overall']['advice'] == sheets[3].cell(grade+2,4).value.split('\n')
        selected += [dto['overall']['diagnosis'],*dto['overall']['advice']]
        missing = [i for i,value in enumerate(selected) if compact(value) not in text]
        assert not missing, f'{file.name}: exact selected text missing {missing}'
        assert '53.75' in text and '50.00' in text
        pages = []
        images = DIRECTORY/'pages'/file.stem
        images.mkdir(parents=True,exist_ok=True)
        for n,p in enumerate(pdf):
            assert p.get_text().strip() and abs(p.rect.width-595.3)<2 and abs(p.rect.height-841.9)<2
            p.get_pixmap(matrix=fitz.Matrix(1,1)).save(images/f'page-{n+1:02d}.png')
            pages.append(len(p.get_text()))
        contact = fitz.open()
        sheet = contact.new_page(width=600,height=849)
        for n,p in enumerate(pdf):
            x,y = (n%3)*200,(n//3)*283
            sheet.insert_image(fitz.Rect(x,y,x+200,y+283),stream=p.get_pixmap(matrix=fitz.Matrix(.5,.5)).tobytes('png'))
        sheet.get_pixmap().save(images/'all-pages.png')
        visible_rings = [d for d in pdf[2].get_drawings() if 180 < d['rect'].y0 and d['rect'].y1 < 390 and d['rect'].width > 30 and d['fill'] is not None and max(d['fill'])-min(d['fill']) > .2]
        evidence={'file':file.name,'sha256':report['fileSha'],'bytes':len(raw),'pages':len(pdf),'overallExact':str(overall),'overallDisplay':str(rounded(overall,2)),'dimensions':13,'modules':4,'rawAnswers':140,'reverseOnce':40,'selectedRuleSegments':len(selected),'missingSegments':missing,'pageTextLengths':pages,'creator':pdf.metadata.get('creator'),'producer':pdf.metadata.get('producer'),'visibleRingVectors':len(visible_rings),'visualSixCharts':'FAIL_ZERO_SCORE_BLANK_RINGS' if len(visible_rings)!=5 else 'VISIBLE_RING_AND_COMPARISON_IMAGE_REVIEWED'}
        output.append(evidence)
        print(json.dumps(evidence,ensure_ascii=True))
    (DIRECTORY/'pdf-oracle.json').write_text(json.dumps(output,indent=2,ensure_ascii=False),'utf-8')
    print('FOUR_REAL_RATIONAL_SQL_EXCEL_PDF_TEXT_PASS')
    if '--require-six-visible' in sys.argv:
        failures = [x['file'] for x in output if x['visibleRingVectors'] != 5]
        assert not failures, f'Six-visible-chart acceptance failed: {failures}; actual zero-score ring regions are blank'

if __name__ == '__main__':
    main()