#!/usr/bin/env python3
"""MT-ZERO-RING-01 frozen customer-text replay: local-only, never run old CLIs.

Required --input-dir contains the four actual Go-rendered frozen DOCX files.
Only a new random result directory is written. Historical FOURBLOCKED remains.
"""
import argparse
from decimal import Decimal
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import subprocess
import sys
import traceback
import uuid
import xml.etree.ElementTree as ET

import openpyxl
import pymupdf as fitz

ROOT = Path(__file__).resolve().parents[2]
OLD = ROOT / 'scripts/test/results/MTRa05eb3af5e49'
OUTPUT = ROOT / 'scripts/test/results/management-traits-zero-ring-frozen-pdf'
NAMES = ['00201-candidate', '00201-tester', '00202-candidate', '00202-tester']
WORKBOOK = ROOT / 'Go-based Refactored System/configs/export-templates/management-traits-002-test-content-v1.xlsx'


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def sha(data):
    return hashlib.sha256(data).hexdigest()


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def scores(dto):
    modules = {x['Key']: x['ChartScore'] for x in dto['modules']}
    return [dto['overall']['chartScore']] + [modules[k] for k in ('self', 'interpersonal', 'task', 'development')]


def opc_contract(oracle, path, dto, original):
    actual = oracle.package(path)
    assert set(actual) == set(original), 'OPC part set changed'
    changed, identical, styles, rings = [], [], [], []
    ring_parts = {f'word/charts/chart{i}.xml' for i in oracle.PARTS}
    for key, number, literal in zip(oracle.KEYS, oracle.PARTS, scores(dto), strict=True):
        part = f'word/charts/chart{number}.xml'
        xml = ET.fromstring(actual[part])
        points = xml.findall('.//c:ser/c:val/c:numLit/c:pt/c:v', oracle.NS)
        value = Decimal(literal)
        assert len(points) == 2 and points[0].text == literal
        assert Decimal(points[1].text) == 100-value
        expected = original[part]
        if value == 0:
            match = re.search(rb'<c:dPt>\s*<c:idx val="1"\s*/>.*?</c:dPt>', expected, re.S)
            assert match is not None
            replacement, count = re.subn(rb'(<c:spPr>\s*)<a:noFill\s*/>',
                rb'\1<a:solidFill><a:srgbClr val="E7E6E6"/></a:solidFill>', match.group(), count=1)
            assert count == 1
            expected = expected[:match.start()] + replacement + expected[match.end():]
        assert oracle.scrub(actual[part], 'c', 'v') == oracle.scrub(expected, 'c', 'v'), part+' unauthorized style change'
        remainder = next(p for p in xml.findall('.//c:ser/c:dPt', oracle.NS) if p.find('c:idx', oracle.NS).get('val') == '1')
        fill = remainder.find('c:spPr', oracle.NS)
        color = fill.find('a:solidFill/a:srgbClr', oracle.NS)
        assert fill.find('a:ln/a:noFill', oracle.NS) is not None
        assert (color is not None and color.get('val') == 'E7E6E6' and fill.find('a:noFill', oracle.NS) is None) if value == 0 else (color is None and fill.find('a:noFill', oracle.NS) is not None)
        rings.append({'key': key, 'exactScore': literal, 'exactZero': value == 0, 'grayOnlyForExactZero': True})
    for part, before in original.items():
        after = actual[part]
        if after == before:
            identical.append(part)
        else:
            changed.append(part)
        if part == 'word/document.xml':
            assert oracle.scrub(after, 'w', 't') == oracle.scrub(before, 'w', 't'), 'non-text document bytes changed'
            styles.append(part)
        elif part == 'word/charts/chart6.xml':
            assert oracle.scrub(after, 'c', 'v') == oracle.scrub(before, 'c', 'v'), 'comparison non-value bytes changed'
            styles.append(part)
        elif part not in ring_parts:
            assert after == before, 'unrelated OPC changed: '+part
    comparison = ET.fromstring(actual['word/charts/chart6.xml']).findall('.//c:ser', oracle.NS)
    assert len(comparison) == 2
    values = [[Decimal(v.text) for v in s.findall('c:val/c:numLit/c:pt/c:v', oracle.NS)] for s in comparison]
    assert values == [[Decimal(d['chartScore']) for d in dto['dimensions']], [Decimal(d['chartNorm']) for d in dto['dimensions']]]
    baseline_norms = ET.fromstring(original['word/charts/chart6.xml']).findall('.//c:ser', oracle.NS)[1]
    assert values[1] == [Decimal(v.text) for v in baseline_norms.findall('c:val/c:numLit/c:pt/c:v', oracle.NS)]
    document_text = ''.join(ET.fromstring(actual['word/document.xml']).itertext())
    assert all(d['score'] in document_text and d['norm'] in document_text for d in dto['dimensions'])
    return {'file': oracle.receipt(path), 'parts': len(actual), 'identicalParts': len(identical),
            'identicalPartNames': identical, 'changedPartNames': changed, 'nonValueDocumentAndComparisonBytesUnchanged': styles,
            'ringStyleBytesUnchangedExceptExactZeroRemainder': True, 'rings': rings,
            'comparisonScores': [str(v) for v in values[0]], 'comparisonNorms': [str(v) for v in values[1]],
            'comparisonDimensions': 13, 'normsUnchanged': True}


def chosen_texts(row, dto, sheets, specs):
    levels = ['excellent', 'good', 'qualified', 'weak', 'insufficient']
    dims = {d['key']: d for d in row['dimensions']}
    selected = []
    assert len(dto['dimensions']) == 13 and len(dto['modules']) == 4 and dto['testOnly']
    for n, d in enumerate(dto['dimensions']):
        assert d['key'] == specs[n][0]
        grade = levels.index(dims[d['key']]['level'])
        for field, sheet, column in [('diagnosis', 1, 4), ('advice', 2, 3)]:
            source = sheets[sheet].cell(n+3, column+grade).value
            assert d[field] == source, 'DTO/XLSX dimension text mismatch'
            selected.append({'key': d['key']+'.'+field, 'text': source})
    for label in ('highest', 'lowest'):
        assert len(dto[label]) == 3
        for item in dto[label]:
            n = [s[0] for s in specs].index(item['Key'])
            source = sheets[0].cell(n+3, 3+levels.index(dims[item['Key']]['level'])).value
            assert item['Text'] == source, 'DTO/XLSX summary mismatch'
            selected.append({'key': label+'.'+item['Key'], 'text': source})
    grade = levels.index(row['run']['level'])
    assert dto['overall']['diagnosis'] == sheets[3].cell(grade+2, 3).value
    assert dto['overall']['advice'] == sheets[3].cell(grade+2, 4).value.split('\n')
    selected.append({'key': 'overall.diagnosis', 'text': dto['overall']['diagnosis']})
    selected.extend({'key': 'overall.advice.'+str(n+1), 'text': t} for n, t in enumerate(dto['overall']['advice']))
    assert len(selected) == 36
    return selected


def extra_pdf(oracle, path, old_path, selected, dto, directory):
    with fitz.open(path) as pdf, fitz.open(old_path) as old:
        assert len(pdf) == 9, 'requires nine physical A4 pages'
        pages = [re.sub(r'\s+', '', p.get_text(sort=False)) for p in pdf]
        full = ''.join(pages)
        assert '\ufffd' not in full and 'TEST' in full and '不可作为人才决策依据' in full
        segments = []
        for item in selected:
            text = re.sub(r'\s+', '', item['text'])
            assert text in full, 'full selected customer text missing: '+item['key']
            segments.append({**item, 'sha256': sha(item['text'].encode()), 'found': True,
                             'wholeSegmentPages': [i+1 for i, p in enumerate(pages) if text in p]})
        for d in dto['dimensions']:
            assert d['name'] in full and d['score'] in full and d['norm'] in full
        assert '53.75' in full and '50.00' in full
        new_spans = oracle.text_signature(pdf[2])
        old_spans = oracle.text_signature(old[2])
        fonts = lambda doc: sorted({(s['font'], round(s['size'], 4)) for p in doc for s in oracle.spans(p)})
        contact = fitz.open()
        sheet = contact.new_page(width=600, height=849)
        for n, page in enumerate(pdf):
            x, y = n%3*200, n//3*283
            sheet.insert_image(fitz.Rect(x, y, x+200, y+283), stream=page.get_pixmap(matrix=fitz.Matrix(.5, .5)).tobytes('png'))
        sheet.get_pixmap().save(directory/'contact.png')
        contact.close()
        pdf[2].get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(directory/'page3-detail.png')
        return {'segments': segments, 'selectedRuleSegments': 36, 'missingSegments': 0,
                'dimensions': 13, 'modules': 4, 'replacementCharacters': 0,
                'crossRenderer': {'oldFile': oracle.receipt(old_path), 'oldProducer': old.metadata.get('producer'),
                    'newProducer': pdf.metadata.get('producer'), 'oldPages': len(old), 'newPages': len(pdf),
                    'oldFonts': fonts(old), 'newFonts': fonts(pdf), 'oldPage3Spans': old_spans, 'newPage3Spans': new_spans,
                    'page3SpansExactEqual': old_spans == new_spans,
                    'coordinateEqualityRequired': False, 'note': 'LO24.2 versus LO26.2; observed positions/fonts retained, not assumed identical'},
                'images': {'page3': str(directory/'page3-detail.png'), 'contact': str(directory/'contact.png')}}


def translated_envelopes(oracle, path, reference, envelopes):
    def slots(pdf):
        labels = [s for s in oracle.spans(pdf[2]) if re.match(r'^\d+\.\d{2} 分', s['text'])
                  and 150 < s['bbox'][1] < 385]
        assert len(labels) == 5, 'five geometry anchor labels required'
        overall = next(s for s in labels if 200 < s['bbox'][0] < 400)
        left = sorted([s for s in labels if s['bbox'][0] < 200], key=lambda s: s['bbox'][1])
        right = sorted([s for s in labels if s['bbox'][0] > 400], key=lambda s: s['bbox'][1])
        assert len(left) == len(right) == 2
        return [overall, *left, *right]
    with fitz.open(path) as actual, fitz.open(reference) as baseline:
        anchor_pairs = list(zip(slots(actual), slots(baseline), strict=True))
        deltas = [a['bbox'][1]-b['bbox'][1] for a, b in anchor_pairs]
        assert max(deltas)-min(deltas) < .002, 'chart anchors did not translate uniformly'
        assert all(a['font'] == b['font'] and a['size'] == b['size'] for a, b in anchor_pairs)
    delta = deltas[0]
    translated = {k: [r[0], r[1]+delta, r[2], r[3]+delta] for k, r in envelopes.items()}
    return translated, {'reference': str(reference), 'fiveAnchorDeltaYpt': deltas,
                        'uniformTranslationYpt': delta, 'shapeAndXUnchanged': True,
                        'reason': 'actual full customer text controls paragraph-relative chart vertical placement'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input-dir', required=True, type=Path)
    parser.add_argument('--soffice', type=Path, default=Path('C:/Program Files/LibreOffice/program/soffice.com'))
    parser.add_argument('--timeout', type=int, default=120)
    args = parser.parse_args()
    assert 1 <= args.timeout <= 180
    directory = OUTPUT/uuid.uuid4().hex[:6]
    directory.mkdir(parents=True, exist_ok=False)
    report = {'task': 'MT-ZERO-RING-01', 'localOnly': True, 'stagingHistoricalStatus': 'FOURBLOCKED',
              'scope': 'unchanged original persisted DTO customer-text fixtures, fresh local LO PDF only',
              'storedRawScores': 'previous independent four-real oracle evidence; not rerun or overwritten',
              'python': sys.version, 'editRounds': 2, 'errors': [], 'cases': []}
    protected = []
    before = {}
    print('FROZEN_RESULTS='+str(directory), flush=True)
    try:
        oracle_path = ROOT/'scripts/test/management-traits-zero-ring-pdf-contract-test.py'
        old_oracle = ROOT/'scripts/test/management-traits-four-real-pdf-oracle.py'
        oracle = load_module('mng_zero_import_only', oracle_path)
        specs = load_module('mng_four_specs_only', old_oracle).SPECS
        protected = [Path(__file__), oracle_path, old_oracle, WORKBOOK, oracle.TEMPLATE]
        protected += [p for p in OLD.rglob('*') if p.is_file()]
        protected += [p for p in args.input_dir.rglob('*') if p.is_file()]
        for root in (ROOT/'scripts/test/results').glob('management-traits-zero-ring-*'):
            if root.resolve() != OUTPUT.resolve():
                protected += [p for p in root.rglob('*') if p.is_file()]
        prior = ROOT/'docs/generated/management-traits-word-candidate-20261001'
        protected += [p for p in prior.rglob('*') if p.is_file()]
        protected = sorted(set(protected))
        before = {str(p): oracle.receipt(p) for p in protected}
        assert oracle.receipt(oracle.TEMPLATE)['sha256'] == oracle.TEMPLATE_SHA
        assert sha(WORKBOOK.read_bytes()) == 'b0498249ae057e3aa1b798922ed2d53a6ca304811943b3b41f4f64f9b024e84c'
        rows = json.loads((OLD/'owned-sql.json').read_text('utf-8'))
        assert len(rows) == 4
        original = oracle.package(oracle.TEMPLATE)
        workbook = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
        dtos, selected = {}, {}
        for name, row in zip(NAMES, rows, strict=True):
            raw = row['report']['dataRaw'].encode('utf-8')
            assert sha(raw) == row['report']['dataSha'], 'original DTO SHA mismatch'
            dto = json.loads(raw)
            assert dto['runId'] == row['run']['id'] and dto['paperId'] == row['paperId'] and dto['examId'] == row['examId']
            dtos[name] = dto
            selected[name] = chosen_texts(row, dto, workbook.worksheets, specs)
            oracle.SCORES[name] = scores(dto)
            old_path = OLD/(name+'.pdf')
            assert sha(old_path.read_bytes()) == row['report']['fileSha'] and old_path.stat().st_size == row['report']['fileBytes']
            report['cases'].append({'name': name, 'dtoSHA256': sha(raw), 'originalDTOUnmodified': True,
                'opc': opc_contract(oracle, args.input_dir/('frozen-'+name+'.docx'), dto, original)})
        workbook.close()
        version = subprocess.run([str(args.soffice), '--version'], capture_output=True, timeout=20, check=True)
        report['libreofficeVersion'] = version.stdout.decode('utf-8', errors='replace').strip()
        assert '26.2.' in report['libreofficeVersion'], 'LO26.2 required'
        pdfs = {}
        for i, name in enumerate(NAMES):
            pdfs[name] = oracle.convert(args.soffice, args.input_dir/('frozen-'+name+'.docx'), directory/str(i), directory/('p'+str(i)), args.timeout)
            print('FROZEN_CONVERT='+name+' '+json.dumps(oracle.receipt(pdfs[name])), flush=True)
        assert all(Decimal(v) == 100 for v in scores(dtos['00202-candidate']))
        envelopes = oracle.geometry(pdfs['00202-candidate'])
        report['geometryReference'] = oracle.receipt(pdfs['00202-candidate'])
        report['geometry'] = envelopes
        for case in report['cases']:
            name = case['name']
            images = directory/('images-'+name)
            case_envelopes, translation = translated_envelopes(oracle, pdfs[name], pdfs['00202-candidate'], envelopes)
            case['geometryTranslation'] = translation
            case['geometry'] = case_envelopes
            evidence = oracle.pdf_contract(pdfs[name], name, case_envelopes, images)
            case['pdf'] = evidence
            case['fullText'] = extra_pdf(oracle, pdfs[name], OLD/(name+'.pdf'), selected[name], dtos[name], images)
            report['errors'].extend(name+': '+e for e in evidence['errors'])
            if name == '00201-tester':
                assert all(Decimal(d['chartScore']) == 0 for d in dtos[name]['dimensions'])
                assert all(Decimal(v) == 0 for v in scores(dtos[name]))
                assert [s['text'].strip() for s in evidence['numericSlots']] == ['0.00 分']*5
                for ring in evidence['rings']:
                    gray = ring['gray']
                    assert gray['rgb'] == [231, 230, 230] and gray['occupiedBins'] == gray['angularBins'] == 360
                    assert gray['whiteHoleOccupancy'] >= .99 and gray['completeVectorCount'] == 1
                assert evidence['visibleCompleteGrayRings'] == 5 and evidence['sixCharts']
            print('FROZEN_GATE='+json.dumps({'name': name, 'pages': evidence['pages'], 'segments': 36,
                  'grayRings': evidence['visibleCompleteGrayRings'], 'errors': evidence['errors']}), flush=True)
        report['selectedRuleSegmentsTotal'] = sum(c['fullText']['selectedRuleSegments'] for c in report['cases'])
        assert report['selectedRuleSegmentsTotal'] == 144
    except Exception as error:
        report['errors'].append(type(error).__name__+': '+str(error))
        report['traceback'] = traceback.format_exc()
    finally:
        if before:
            after = {str(p): {'path': str(p), 'bytes': p.stat().st_size, 'sha256': sha(p.read_bytes())} for p in protected}
            report['protectedFiles'] = after
            report['protectedUnchanged'] = before == after
            if before != after:
                report['errors'].append('protected original assets changed')
        report['status'] = 'FAIL' if report['errors'] else 'LOCAL_PASS'
        report['outputCode'] = 1 if report['errors'] else 0
        save(directory/'contract.json', report)
        print('FROZEN_CONTRACT='+json.dumps({'status': report['status'], 'exit': report['outputCode'],
              'errors': report['errors'], 'evidence': str(directory/'contract.json'),
              'sha256': sha((directory/'contract.json').read_bytes())}), flush=True)
    return report['outputCode']


if __name__ == '__main__':
    sys.exit(main())