"""Independent local oracle for the two synthetic 005 customer-template PDFs."""
import base64
import hashlib
import importlib.util
import json
from decimal import Decimal
from fractions import Fraction
from pathlib import Path
import re
import sys

import openpyxl
import pymupdf as fitz

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = Path(sys.argv[sys.argv.index("--output") + 1]).resolve() if "--output" in sys.argv else ROOT / "scripts/test/results/mng005-report-e2e-20261008"
E2E_SOURCE = Path(__file__).with_name("management-traits-005-report-e2e-local-20261008.js")
CANONICAL_SOURCE = ROOT / "Go-based Refactored System/internal/service/management_traits_identity.go"
CANONICAL_CONTRACT = OUTPUT / "canonical-contract-v2.json"
IDENTITY_CONTRACT = OUTPUT / "identity-commitment.json"
IDENTITY_CONTRACT_V2 = OUTPUT / ("identity-commitment-v2.json" if (OUTPUT / "identity-commitment-v2.json").exists() else "identity-commitment-v2-blocked.json")
SAFE_RECEIPTS = (
    OUTPUT / "browser-receipt.json",
    OUTPUT / "db-file-audit-before-cleanup.json",
    OUTPUT / "cleanup-final.json",
)
IDENTITY_LABELS = ("姓名：", "性别：", "所在单位：", "职务：", "联系方式：", "测评日期：")
JWT_LIKE = re.compile(r"(?<![A-Za-z0-9_-])eyJ[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+(?![A-Za-z0-9_-])")
RAW_EXAM_ID = re.compile(r"(?<!\d)\d{19}(?!\d)")
RAW_UUID = re.compile(r"(?i)(?<![a-f0-9])[a-f0-9]{8}-[a-f0-9]{4}-[1-5][a-f0-9]{3}-[89ab][a-f0-9]{3}-[a-f0-9]{12}(?![a-f0-9])")
PHONE = re.compile(r"(?<!\d)1\d{10}(?!\d)")
SENSITIVE_KEY_PARTS = ("password", "passwd", "pwd", "secret", "token", "authorization", "cookie", "dsn", "credential")
CREDENTIAL_VALUE = re.compile(
    r"-----BEGIN [A-Z ]*PRIVATE KEY-----|\bBearer\s+[A-Za-z0-9._~+/-]+=*|"
    r"(?:mysql|postgres(?:ql)?|redis|mongodb(?:\+srv)?)://[^\s/:@]+:[^\s@]+@|"
    r"(?:password|passwd|pwd|secret|token|authorization)\s*[:=]\s*[^\s,;]{4,}",
    re.IGNORECASE,
)
BASE64_VALUE = re.compile(r"[A-Za-z0-9+/]+={0,2}")
REDACTED_VALUES = {"", "***", "redacted", "<redacted>", "[redacted]", "none", "null"}
MAX_DECODED_BYTES = 64 * 1024
MAX_SINGLE_PDF_OBJECT_BYTES = 16 * 1024 * 1024
MAX_TOTAL_PDF_OBJECT_BYTES = 128 * 1024 * 1024
MAX_SINGLE_ATTACHMENT_BYTES = 16 * 1024 * 1024
MAX_TOTAL_ATTACHMENT_BYTES = 64 * 1024 * 1024
spec = importlib.util.spec_from_file_location("mng_existing_oracle", Path(__file__).with_name("management-traits-four-real-pdf-oracle.py"))
assert spec and spec.loader
existing = importlib.util.module_from_spec(spec)
spec.loader.exec_module(existing)


def digest(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def compact_label(value: str) -> str:
    return re.sub(r"\s+", "", value)


def extract_identity(lines: list[str], label: str) -> str:
    index = next(i for i, value in enumerate(lines) if compact_label(value) == label)
    if index + 1 == len(lines) or compact_label(lines[index + 1]) in IDENTITY_LABELS:
        return ""
    return lines[index + 1]


def hash_parts(parts: list[str]) -> str:
    return digest("\0".join(parts).encode("utf-8"))


def scan_string(value: str, location: str, findings: list[str], key_hint: str = "", depth: int = 0) -> None:
    if JWT_LIKE.search(value) or CREDENTIAL_VALUE.search(value):
        findings.append(location)
    compact = value.strip()
    base64_hint = bool(re.search(r"base64|payload|blob|content|data", key_hint, re.IGNORECASE))
    candidate = 16 <= len(compact) <= 8192 and len(compact) % 4 == 0 and BASE64_VALUE.fullmatch(compact)
    if depth >= 2 or not candidate or not (base64_hint or "=" in compact or "+" in compact or "/" in compact):
        return
    try:
        decoded = base64.b64decode(compact, validate=True)
    except ValueError:
        return
    if not decoded or len(decoded) > MAX_DECODED_BYTES:
        return
    try:
        text = decoded.decode("utf-8")
    except UnicodeDecodeError:
        return
    printable = sum(character.isprintable() or character in "\r\n\t" for character in text)
    if printable / len(text) < 0.85:
        return
    try:
        nested = json.loads(text)
    except json.JSONDecodeError:
        scan_string(text, f"{location}:base64", findings, depth=depth + 1)
    else:
        scan_json(nested, findings, f"{location}:base64", depth + 1)


def scan_json(value: object, findings: list[str], location: str = "$", depth: int = 0) -> None:
    if isinstance(value, dict):
        for key, child in value.items():
            child_location = f"{location}.{key}"
            normalized_key = re.sub(r"[^a-z0-9]", "", str(key).lower())
            if any(part in normalized_key for part in SENSITIVE_KEY_PARTS) and child is not None and str(child).strip().lower() not in REDACTED_VALUES:
                findings.append(child_location)
            if isinstance(child, str):
                if RAW_EXAM_ID.search(child) or RAW_UUID.search(child):
                    findings.append(child_location)
                scan_string(child, child_location, findings, str(key), depth)
            else:
                scan_json(child, findings, child_location, depth)
    elif isinstance(value, list):
        for index, child in enumerate(value):
            scan_json(child, findings, f"{location}[{index}]", depth)
    elif isinstance(value, str):
        if RAW_EXAM_ID.search(value) or RAW_UUID.search(value):
            findings.append(location)
        scan_string(value, location, findings, depth=depth)


def validate_scanner_contract() -> None:
    findings: list[str] = []
    scan_json(json.loads(r'{"pa\u0073sword":"synthetic-value"}'), findings)
    assert findings == ["$.password"]
    nested = base64.b64encode(b'{"authToken":"synthetic-value"}').decode("ascii")
    findings = []
    scan_json({"payloadBase64": nested}, findings)
    assert findings == ["$.payloadBase64:base64.authToken"]
    findings = []
    scan_json({"password": "***", "token": "redacted"}, findings)
    assert not findings


def scan_bytes(value: bytes, location: str, findings: list[str]) -> None:
    if JWT_LIKE.search(value.decode("latin1")) or CREDENTIAL_VALUE.search(value.decode("latin1")):
        findings.append(location)
    for encoding in ("utf-8", "utf-16-le", "utf-16-be", "latin1"):
        try:
            text = value.decode(encoding)
        except UnicodeDecodeError:
            continue
        scan_string(text, f"{location}:{encoding}", findings)


def enforce_pdf_object_budget(object_text_bytes: int, raw_stream_bytes: int, decoded_stream_bytes: int, raw_size: int, decoded_size: int) -> tuple[int, int]:
    if raw_size > MAX_SINGLE_PDF_OBJECT_BYTES or decoded_size > MAX_SINGLE_PDF_OBJECT_BYTES:
        raise AssertionError("PDF single object budget exceeded")
    raw_stream_bytes += raw_size
    decoded_stream_bytes += decoded_size
    if object_text_bytes + raw_stream_bytes + decoded_stream_bytes > MAX_TOTAL_PDF_OBJECT_BYTES:
        raise AssertionError("PDF total object budget exceeded")
    return raw_stream_bytes, decoded_stream_bytes


def validate_budget_contract() -> None:
    legal = 10_688_445
    raw, decoded = enforce_pdf_object_budget(1024, 0, 0, 4096, legal)
    assert (raw, decoded) == (4096, legal)
    single_rejected = total_rejected = False
    try:
        enforce_pdf_object_budget(0, 0, 0, 1, MAX_SINGLE_PDF_OBJECT_BYTES + 1)
    except AssertionError:
        single_rejected = True
    try:
        enforce_pdf_object_budget(MAX_TOTAL_PDF_OBJECT_BYTES - 10, 0, 0, 6, 5)
    except AssertionError:
        total_rejected = True
    assert single_rejected and total_rejected
    print(json.dumps({"status": "MNG005_PDF_BUDGET_CONTRACT_PASS", "maxSingleObjectBytes": MAX_SINGLE_PDF_OBJECT_BYTES, "maxTotalObjectBytes": MAX_TOTAL_PDF_OBJECT_BYTES, "legalDecodedStreamBytes": legal, "singleOverflowRejected": single_rejected, "totalOverflowRejected": total_rejected}))


def scan_pdf_credentials(pdf: fitz.Document, texts: list[str]) -> dict[str, int]:
    assert not pdf.needs_pass and not pdf.is_encrypted
    findings: list[str] = []
    for index, text in enumerate(texts):
        scan_string(text, f"page[{index}]", findings)
    metadata = pdf.metadata or {}
    scan_json(metadata, findings, "metadata")
    trailer = pdf.pdf_trailer()
    scan_string(trailer, "trailer", findings)
    catalog_xref = pdf.pdf_catalog()
    scan_string(pdf.xref_object(catalog_xref, compressed=False), "catalog", findings)
    xref_objects = 0
    stream_objects = 0
    object_text_bytes = 0
    raw_stream_bytes = 0
    decoded_stream_bytes = 0
    for xref in range(1, pdf.xref_length()):
        try:
            value = pdf.xref_object(xref, compressed=False)
        except RuntimeError:
            if pdf.xref_is_stream(xref):
                raise AssertionError(f"unreadable stream object {xref}")
            continue
        xref_objects += 1
        object_text_bytes += len(value.encode("utf-8"))
        assert object_text_bytes <= MAX_TOTAL_PDF_OBJECT_BYTES
        scan_string(value, f"xref[{xref}]", findings)
        if not pdf.xref_is_stream(xref):
            continue
        stream_objects += 1
        try:
            raw_stream = pdf.xref_stream_raw(xref)
            decoded_stream = pdf.xref_stream(xref)
        except RuntimeError as error:
            raise AssertionError(f"unreadable stream object {xref}") from error
        assert raw_stream is not None and decoded_stream is not None
        raw_stream_bytes, decoded_stream_bytes = enforce_pdf_object_budget(object_text_bytes, raw_stream_bytes, decoded_stream_bytes, len(raw_stream), len(decoded_stream))
        scan_bytes(raw_stream, f"xref[{xref}]:raw", findings)
        scan_bytes(decoded_stream, f"xref[{xref}]:decoded", findings)
    attachment_names = list(pdf.embfile_names())
    attachment_bytes = 0
    for name in attachment_names:
        scan_string(name, "attachment-name", findings)
        payload = pdf.embfile_get(name)
        assert len(payload) <= MAX_SINGLE_ATTACHMENT_BYTES
        attachment_bytes += len(payload)
        assert attachment_bytes <= MAX_TOTAL_ATTACHMENT_BYTES
        scan_bytes(payload, f"attachment[{name}]", findings)
    assert not findings
    return {
        "findings": 0,
        "needsPass": False,
        "encrypted": False,
        "metadataValues": len(metadata),
        "trailerScanned": 1,
        "catalogScanned": 1,
        "xrefObjects": xref_objects,
        "streamObjects": stream_objects,
        "objectTextBytes": object_text_bytes,
        "rawStreamBytes": raw_stream_bytes,
        "decodedStreamBytes": decoded_stream_bytes,
        "attachments": len(attachment_names),
        "attachmentBytes": attachment_bytes,
    }


def main() -> None:
    validate_scanner_contract()
    core_name = "e2e-core-receipt.json" if (OUTPUT / "e2e-core-receipt.json").exists() else "report-only-receipt.json"
    complete_v2 = (OUTPUT / core_name).exists()
    receipt_path = OUTPUT / (core_name if complete_v2 else "browser-receipt.json")
    receipt_bytes = receipt_path.read_bytes()
    receipt = json.loads(receipt_bytes)
    assert receipt["schema"] in ({"mng005-customer-report-e2e-core-v2", "mng005-report-only-v1", "mng005-report-only-v2", "mng005-report-only-v3", "mng005-report-only-v4"} if complete_v2 else {"mng005-customer-report-e2e-v1"})
    if receipt["schema"] in {"mng005-report-only-v2", "mng005-report-only-v3", "mng005-report-only-v4"}:
        assert receipt["answerClicks"] == receipt["answerSaveRequests"] == receipt["startRequests"] == receipt["submitRequests"] == 0
        assert re.fullmatch(r"\d{13}", receipt.get("generation", receipt.get("baselineGeneration", "")))
    canonical_bytes = CANONICAL_CONTRACT.read_bytes()
    canonical = json.loads(canonical_bytes)
    assert canonical["schema"] == "mng005-canonical-contract-v2"
    assert digest(CANONICAL_SOURCE.read_bytes()) == canonical["sourceSHA256"]
    for source in canonical["sources"]:
        assert digest((ROOT / source["path"]).read_bytes()) == source["sha256"]
    assert set(canonical["astNodes"]) == {"dimensionCatalog", "dimensionBuilder", "scoringFunction", "moduleAggregation"}
    assert all(re.fullmatch(r"[a-f0-9]{64}", node["sha256"]) for node in canonical["astNodes"].values())
    assert len(canonical["dimensions"]) == 13 and len(canonical["modules"]) == 4
    identity_bytes = IDENTITY_CONTRACT.read_bytes() if not complete_v2 else b""
    identity_contract = json.loads(identity_bytes) if identity_bytes else None
    if not complete_v2:
        assert identity_contract["schema"] == "mng005-identity-commitment-v1"
        assert digest(receipt_bytes) == identity_contract["browserReceiptSHA256"]
    identity_v2_bytes = IDENTITY_CONTRACT_V2.read_bytes()
    identity_v2 = json.loads(identity_v2_bytes)
    assert identity_v2["schema"] == "mng005-identity-commitment-v2"
    if complete_v2:
        assert identity_v2["status"] == "COMPLETE"
        assert identity_v2["coreReceiptSHA256"] == digest(receipt_bytes)
    else:
        assert identity_v2["status"] == "BLOCKED_MISSING_IMMUTABLE_SOURCE_BINDING"
        assert identity_v2["missing"] == ["00501:runIdHash", "00502:runIdHash"] and identity_v2["records"] == []
    e2e_source = E2E_SOURCE.read_text(encoding="utf-8")
    begin_marker = "// IDENTITY_COMMITMENT_FORMULA_BEGIN"
    end_marker = "// IDENTITY_COMMITMENT_FORMULA_END"
    begin, end = e2e_source.index(begin_marker), e2e_source.index(end_marker)
    formula_bytes = e2e_source[begin + len(begin_marker):end].encode("utf-8")
    assert digest(formula_bytes) == identity_v2["formulaSourceSHA256"]
    assert digest(E2E_SOURCE.read_bytes()) == identity_v2["producerSourceSHA256"]
    if complete_v2:
        assert identity_v2["coreReceiptSHA256"] == digest(receipt_bytes)
    else:
        assert digest(receipt_bytes) == identity_v2["safeReceiptSHA256"]
    receipt_findings: list[str] = []
    safe_receipts = tuple(path for path in (receipt_path, IDENTITY_CONTRACT_V2, OUTPUT / "cleanup-receipt.json") if path.exists()) if complete_v2 else SAFE_RECEIPTS
    for path in safe_receipts:
        scan_json(json.loads(path.read_text(encoding="utf-8")), receipt_findings, path.name)
    assert not receipt_findings
    workbook = ROOT / "docs/260929管理特质测评-优化/260928测评内容+数据图.xlsx"
    assert digest(workbook.read_bytes()) == "b0498249ae057e3aa1b798922ed2d53a6ca304811943b3b41f4f64f9b024e84c"
    sheets = openpyxl.load_workbook(workbook, data_only=True, read_only=True).worksheets
    selected = []
    for index, (_, _, _, _) in enumerate(existing.SPECS):
        selected.extend([sheets[1].cell(index + 3, 6).value, sheets[2].cell(index + 3, 5).value])
    dimension_names = [sheets[1].cell(index + 3, 2).value for index in range(13)]
    assert all(isinstance(value, str) and value.strip() for value in dimension_names)
    for index in list(range(3)) + list(range(3)):
        selected.append(sheets[0].cell(index + 3, 5).value)
    selected.extend([sheets[3].cell(4, 3).value, *sheets[3].cell(4, 4).value.split("\n")])
    assert len(selected) == 36 and all(isinstance(value, str) and value.strip() for value in selected)
    independent_dimensions: dict[str, Fraction] = {}
    reverse_count = 0
    for dimension in canonical["dimensions"]:
        finals = []
        for item in dimension["items"]:
            raw = 3
            finals.append(6 - raw if item["reverse"] else raw)
            reverse_count += item["reverse"]
        score = existing.Fraction(25 * (sum(finals) - len(finals)), len(finals))
        assert score == 50
        independent_dimensions[dimension["key"]] = score
    independent_overall = sum(independent_dimensions.values(), existing.Fraction()) / 13
    assert independent_overall == 50 and reverse_count == 40
    expected_module_keys = tuple(module["key"] for module in canonical["modules"])
    assert expected_module_keys == ("self", "interpersonal", "task", "development")
    independent_modules = {}
    for module in canonical["modules"]:
        children = [independent_dimensions[key] for key in module["dimensions"]]
        independent_modules[module["key"]] = sum(children, Fraction()) / len(children)
    assert list(independent_modules) == list(expected_module_keys)
    assert len(independent_modules) == len(set(independent_modules)) == 4
    assert all(score == 50 for score in independent_modules.values())

    results = []
    by_code = {record["code"]: record for record in receipt["records"]}
    identity_by_code = {record.get("code", record.get("productCode")): record for record in (identity_v2["records"] if complete_v2 else identity_contract["records"])}
    assert set(by_code) == set(identity_by_code) == {"00501", "00502"}
    for code in ("00501", "00502"):
        record = by_code[code]
        identity = identity_by_code[code]
        pdf_name = record.get("file", record.get("pdf"))
        expected_bytes = record.get("bytes", record.get("pdfBytes"))
        expected_sha = record.get("sha256", record.get("pdfSHA256"))
        path = OUTPUT / pdf_name
        raw = path.read_bytes()
        assert raw.startswith(b"%PDF-")
        assert len(raw) == expected_bytes and digest(raw) == expected_sha
        pdf = fitz.open(path)
        texts = [page.get_text(sort=False) for page in pdf]
        compact = existing.compact("".join(texts))
        lines = [line.strip() for text in texts for line in text.splitlines() if line.strip()]
        extracted_identity = {
            "name": extract_identity(lines, "姓名："),
            "gender": extract_identity(lines, "性别："),
            "affiliation": extract_identity(lines, "所在单位："),
            "post": extract_identity(lines, "职务："),
            "phone": extract_identity(lines, "联系方式："),
        }
        assert extracted_identity["affiliation"] == extracted_identity["post"] == ""
        extracted_field_hashes = {
            field: hash_parts([(identity_v2 if complete_v2 else identity_contract)["identityFieldDomain"], code, field, extracted_identity[field]])
            for field in ("name", "gender", "phone")
        }
        assert extracted_field_hashes == identity["identityFieldsSHA"]
        if complete_v2:
            core = by_code[code]
            assert identity["examIdHash"] == core.get("examIdHash", digest(core["examId"].encode()) if "examId" in core else "")
            assert identity["runIdHash"] == core["runIdHash"] and identity["reportIdHash"] == core["reportIdHash"]
            assert identity["titleHash"] == core["titleHash"] and identity["coreReceiptSHA256"] == digest(receipt_bytes)
            assert identity["finalCommitment"] == hash_parts([identity_v2["identityDomain"], identity_v2["formulaVersion"], identity_v2["formulaSourceSHA256"], code, identity["examIdHash"], identity["runIdHash"], identity["reportIdHash"], identity["titleHash"], identity_v2["producerSourceSHA256"], identity_v2["coreReceiptSHA256"], identity["pdfSHA256"], extracted_field_hashes["name"], extracted_field_hashes["gender"], extracted_field_hashes["phone"]])
        else:
            assert identity["identityCommitment"] == hash_parts([identity_contract["identityDomain"], identity["examIdHash"], code, extracted_field_hashes["name"], extracted_field_hashes["gender"], extracted_field_hashes["phone"]])
        assert identity["pdf"] == pdf_name and identity["pdfBytes"] == len(raw) and identity["pdfSHA256"] == digest(raw)
        assert identity["reportIdHash"] == record["reportIdHash"]
        phones = set(PHONE.findall("".join(texts)))
        assert len(phones) == 1
        assert hash_parts([(identity_v2 if complete_v2 else identity_contract)["identityFieldDomain"], code, "phone", next(iter(phones))]) == identity["identityFieldsSHA"]["phone"]
        pdf_scan = scan_pdf_credentials(pdf, texts)
        assert len(pdf) == 9
        assert all(text.strip() for text in texts)
        assert all(abs(page.rect.width - 595.3) < 2 and abs(page.rect.height - 841.9) < 2 for page in pdf)
        assert "TEST" in compact and "不可作为人才决策依据" in compact
        star_glyphs = sum(text.count("★") + text.count("☆") for text in texts)
        assert star_glyphs == 4
        assert "\ufffd" not in compact
        assert not any(token in compact for token in ("report.testLabel", "report.testTitle", "dimension.", "{{", "}}"))
        assert all(existing.compact(value) in compact for value in selected)
        assert all(existing.compact(value) in compact for value in dimension_names)
        assert all(not page.get_links() for page in pdf)
        drawings = pdf[2].get_drawings()
        rings = [item for item in drawings if item["fill"] and 150 < item["rect"].y0 < 300 and 70 < item["rect"].height < 150 and 35 < item["rect"].width < 75 and len(item["items"]) == 22]
        bars = [item for item in drawings if item["fill"] and 440 < item["rect"].y0 < 465 and 525 < item["rect"].y1 < 535 and 8 < item["rect"].width < 11 and len(item["items"]) == 4]
        lines = [item for item in drawings if item["type"] == "s" and 430 < item["rect"].y0 < 455 and item["rect"].width > 350 and len(item["items"]) == 12]
        assert (len(rings), len(bars), len(lines)) == (5, 13, 1)
        if not complete_v2:
            assert Decimal(record["expected"]) == Decimal("50.000000") and Decimal(record["actual"]) == Decimal("50")
        results.append({"code": code, "rawPattern": "140x3", "reverseMappedItems": reverse_count, "independentDimensionScores": 13, "independentModuleScores": {key: str(value) for key, value in independent_modules.items()}, "independentOverallExact": str(independent_overall), "bytes": len(raw), "sha256": digest(raw), "pages": len(pdf), "a4Pages": len(pdf), "nonemptyPages": len(pdf), "customerTextSegments": 36, "dimensionNames": 13, "independentModuleFacts": len(independent_modules), "ringCharts": len(rings), "comparisonBars": len(bars), "normPolylines": len(lines), "starGlyphs": star_glyphs, "externalLinks": 0, "testLabel": True, "placeholders": 0, "privacyVerdict": "PASS_SYNTHETIC_IDENTITY_COMMITTED", "identityCommitment": identity.get("finalCommitment", identity.get("identityCommitment")), "runIdHashPresent": bool(identity.get("runIdHash")), "identityFieldMatches": 3, "phoneValues": len(phones), "credentialFindings": 0, "pdfObjectCredentialScan": pdf_scan})
    verdict = {"status": "PASS_LOCAL_SYNTHETIC_ONLY_WITH_IDENTITY_V2_COMPLETE" if complete_v2 else "PASS_LOCAL_SYNTHETIC_ONLY_WITH_IDENTITY_BINDING_BLOCKED", "identityBindingBlocker": [] if complete_v2 else identity_v2["missing"], "canonicalSources": canonical["sources"], "canonicalASTNodeHashes": {key: value["sha256"] for key, value in canonical["astNodes"].items()}, "canonicalSourceSHA": canonical["sourceSHA256"], "canonicalFunctionSHA": canonical["functionSHA256"], "canonicalContractSHA": digest(canonical_bytes), "legacyIdentityCommitmentSHA": digest(identity_bytes) if identity_bytes else None, "identityFormulaSourceSHA": identity_v2["formulaSourceSHA256"], "identityV2ReceiptSHA": digest(identity_v2_bytes), "objectBudgets": {"maxSingleObjectBytes": MAX_SINGLE_PDF_OBJECT_BYTES, "maxTotalObjectBytes": MAX_TOTAL_PDF_OBJECT_BYTES}, "recursiveReceiptCredentialScan": {"files": len(safe_receipts), "findings": 0, "maxBase64Depth": 2, "maxDecodedBytes": MAX_DECODED_BYTES}, "pdfObjectCredentialScan": {"documents": len(results), "findings": sum(record["pdfObjectCredentialScan"]["findings"] for record in results), "encryptedDocuments": sum(int(record["pdfObjectCredentialScan"]["encrypted"]) for record in results), "needsPassDocuments": sum(int(record["pdfObjectCredentialScan"]["needsPass"]) for record in results), "xrefObjects": sum(record["pdfObjectCredentialScan"]["xrefObjects"] for record in results), "streamObjects": sum(record["pdfObjectCredentialScan"]["streamObjects"] for record in results), "objectTextBytes": sum(record["pdfObjectCredentialScan"]["objectTextBytes"] for record in results), "rawStreamBytes": sum(record["pdfObjectCredentialScan"]["rawStreamBytes"] for record in results), "decodedStreamBytes": sum(record["pdfObjectCredentialScan"]["decodedStreamBytes"] for record in results), "attachments": sum(record["pdfObjectCredentialScan"]["attachments"] for record in results), "attachmentBytes": sum(record["pdfObjectCredentialScan"]["attachmentBytes"] for record in results)}, "records": results}
    (OUTPUT / "pdf-oracle.json").write_text(json.dumps(verdict, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(verdict, ensure_ascii=False))


if __name__ == "__main__":
    if "--budget-contract" in sys.argv:
        validate_budget_contract()
    else:
        main()
