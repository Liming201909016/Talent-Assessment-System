"""Offline independent Fraction/XLSX/OPC/PDF oracle; synthetic fixture only."""
import hashlib
import importlib.util
import json
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
import pymupdf as fitz

ROOT = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location("mng_existing_oracle", Path(__file__).with_name("management-traits-four-real-pdf-oracle.py"))
assert spec and spec.loader
existing = importlib.util.module_from_spec(spec)
spec.loader.exec_module(existing)


def main():
    directory = Path(sys.argv[1]).resolve()
    assert directory.parent == ROOT / "scripts/test/results"
    read = lambda name: (directory / name).read_bytes()
    sha = lambda raw: hashlib.sha256(raw).hexdigest()
    source_raw, snapshot_raw = read("source.json"), read("snapshot.json")
    source, snapshot = json.loads(source_raw), json.loads(snapshot_raw)
    receipt = json.loads(read("render-receipt.json"))
    assert source["synthetic"] and snapshot["synthetic"] and receipt["synthetic"]
    assert source["policy"] == "legacy_verified_snapshot" and snapshot["policy"] == "audited_reissue"
    assert sha(source_raw) == snapshot["sourceSha"] == receipt["sourceSHA"]
    assert sha(snapshot_raw) == receipt["snapshotSHA"]
    assert sha(read("audit-signature.bin")) == snapshot["signatureSha"]
    assert sha(read("synthetic-auditor-public.bin")) == snapshot["auditorKeySha"]
    assert (datetime.fromisoformat(source["owner"]["submittedAt"]) - datetime.fromisoformat(source["owner"]["startedAt"])).total_seconds() == snapshot["userTimeSeconds"] == 720
    dto = snapshot["report"]
    assert dto["testOnly"] and dto["participant"]["name"] == "SYNTHETIC REISSUE SAMPLE"
    assert dto["submittedAt"] == "2026-09-01 10:12"
    answers = {a["number"]: a for a in source["input"]["answers"]}
    manifest = {q["number"]: q for q in source["manifest"]["questions"]}
    mapping = {q["number"]: q for q in source["mapping"]["questions"]}
    buckets = {b["paperQuestionId"]: b for b in source["buckets"]}
    assert sorted(answers) == sorted(manifest) == sorted(mapping) == list(range(1, 141))
    assert len(buckets) == 140
    scores, norms, evidence = {}, {}, {d["key"]: d for d in snapshot["scores"]}
    reverse_count = 0
    for key, module, terms, norm in existing.SPECS:
        total = 0
        for term in terms:
            number = abs(term)
            a, m, q = answers[number], manifest[number], mapping[number]
            assert a["answered"] and 1 <= a["raw"] <= 5
            assert m["dimensionKey"] == key and m["reverse"] == (term < 0)
            assert a["sourceQuestionId"] == q["sourceQuestionId"]
            assert next(o["raw"] for o in q["options"] if o["sourceOptionId"] == a["selectedOptionId"]) == a["raw"]
            bucket = buckets[a["paperQuestionId"]]
            assert bucket["sourceQuestionId"] == a["sourceQuestionId"]
            assert sorted(o["raw"] for o in bucket["options"]) == [1, 2, 3, 4, 5]
            assert [o["raw"] for o in bucket["options"] if o["checked"]] == [a["raw"]]
            assert [o["raw"] for o in bucket["options"] if o["isRight"]] == [1 if term < 0 else 5]
            final = 6 - a["raw"] if term < 0 else a["raw"]
            assert final == 3
            reverse_count += term < 0
            total += final
        value = existing.Fraction(25 * (total - len(terms)), len(terms))
        d = evidence[key]
        assert (d["sum"], d["count"], existing.Fraction(d["exact"]), existing.Fraction(d["norm"]), d["level"]) == (total, len(terms), value, norm, existing.level(value))
        scores[key], norms[key] = value, norm
    overall = sum(scores.values(), existing.Fraction()) / 13
    assert existing.Fraction(snapshot["overallExact"]) == overall == 50
    assert existing.Fraction(snapshot["overallNormExact"]) == existing.Fraction(705, 13)
    assert reverse_count == snapshot["reverseCount"] == 40
    for m in dto["modules"]:
        children = [scores[key] for key, module, _, _ in existing.SPECS if module == m["Key"]]
        value = sum(children, existing.Fraction()) / len(children)
        assert Decimal(m["ChartScore"]) == existing.rounded(value, 12)
        stored = next(x for x in snapshot["modules"] if x["key"] == m["Key"])
        assert stored["count"] == len(children) and existing.Fraction(stored["exact"]) == value
    assert len(dto["dimensions"]) == 13 and len(dto["modules"]) == 4
    assert [s["Key"] for s in dto["highest"]] == [s[0] for s in existing.SPECS[:3]] == [s["Key"] for s in dto["lowest"]]
    workbook = ROOT / "docs/260929管理特质测评-优化/260928测评内容+数据图.xlsx"
    assert sha(workbook.read_bytes()) == receipt["contentSHA"] == dto["contentSourceSha"] == "b0498249ae057e3aa1b798922ed2d53a6ca304811943b3b41f4f64f9b024e84c"
    sheets = openpyxl.load_workbook(workbook, data_only=True).worksheets
    selected = []
    for i, d in enumerate(dto["dimensions"]):
        assert d["key"] == existing.SPECS[i][0] and d["score"] == "50.00" and d["level"] == "合格"
        assert d["norm"] == str(existing.rounded(norms[d["key"]], 2))
        assert d["diagnosis"] == sheets[1].cell(i + 3, 6).value
        assert d["advice"] == sheets[2].cell(i + 3, 5).value
        selected += [d["diagnosis"], d["advice"]]
    for s in dto["highest"] + dto["lowest"]:
        i = [r[0] for r in existing.SPECS].index(s["Key"])
        assert s["Text"] == sheets[0].cell(i + 3, 5).value
        selected.append(s["Text"])
    assert dto["overall"]["diagnosis"] == sheets[3].cell(4, 3).value
    assert dto["overall"]["advice"] == sheets[3].cell(4, 4).value.split("\n")
    selected += [dto["overall"]["diagnosis"], *dto["overall"]["advice"]]
    assert len(selected) == 36
    pdf_raw = read("report.pdf")
    assert pdf_raw.startswith(b"%PDF-") and sha(pdf_raw) == receipt["pdfSHA"] and len(pdf_raw) == receipt["pdfBytes"]
    assert sha(read("report.docx")) == receipt["docxSHA"]
    pdf = fitz.open(directory / "report.pdf")
    text = existing.compact("".join(p.get_text(sort=False) for p in pdf))
    assert "TEST" in text and "不可作为人才决策依据" in text and "SYNTHETICREISSUESAMPLE" in text
    assert "\ufffd" not in text and "report.testLabel" not in text
    assert all(existing.compact(s) in text for s in selected), "full customer text missing"
    assert all(existing.compact(d["name"]) in text for d in dto["dimensions"])
    for page in pdf:
        assert page.get_text().strip() and abs(page.rect.width - 595.3) < 2 and abs(page.rect.height - 841.9) < 2
    drawings = pdf[2].get_drawings()
    rings = [d for d in drawings if d["fill"] and 150 < d["rect"].y0 < 300 and 70 < d["rect"].height < 150 and 35 < d["rect"].width < 75 and len(d["items"]) == 22]
    bars = [d for d in drawings if d["fill"] and 440 < d["rect"].y0 < 465 and 525 < d["rect"].y1 < 535 and 8 < d["rect"].width < 11 and len(d["items"]) == 4]
    lines = [d for d in drawings if d["type"] == "s" and 430 < d["rect"].y0 < 455 and d["rect"].width > 350 and len(d["items"]) == 12]
    assert (len(rings), len(bars), len(lines)) == (5, 13, 1), "actual six-chart geometry missing"
    with zipfile.ZipFile(directory / "report.docx") as z:
        ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
        for index in range(1, 6):
            values = ET.fromstring(z.read(f"word/charts/chart{index}.xml")).findall(".//c:val/c:numLit/c:pt/c:v", ns)
            assert [Decimal(v.text) for v in values] == [Decimal(50), Decimal(50)]
        values = ET.fromstring(z.read("word/charts/chart6.xml")).findall(".//c:val/c:numLit/c:pt/c:v", ns)
        assert [Decimal(v.text) for v in values[:13]] == [Decimal(50)] * 13
        assert [Decimal(v.text) for v in values[13:]] == [existing.rounded(s[3], 12) for s in existing.SPECS]
    originals = ROOT / "docs/260929管理特质测评-优化"
    before = json.loads(read("source-start.json"))
    original_evidence = []
    for file in originals.iterdir():
        assert sha(file.read_bytes()) == before[file.relative_to(ROOT).as_posix()]
        with zipfile.ZipFile(file) as z:
            styles = [n for n in z.namelist() if "styles" in n.lower() and n.endswith(".xml")]
            original_evidence.append({"file": file.name, "sha": sha(file.read_bytes()), "styles": {n: sha(z.read(n)) for n in styles}})
    pdf[2].get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(directory / "page-03.png")
    contact = fitz.open()
    sheet = contact.new_page(width=600, height=849)
    for n, page in enumerate(pdf):
        x, y = (n % 3) * 200, (n // 3) * 283
        sheet.insert_image(fitz.Rect(x, y, x + 200, y + 283), stream=page.get_pixmap(matrix=fitz.Matrix(.5, .5)).tobytes("png"))
    sheet.get_pixmap().save(directory / "all-pages.png")
    result = {"status": "PASS_LOCAL_SYNTHETIC_ONLY", "pages": len(pdf), "bytes": len(pdf_raw), "pdfSHA": sha(pdf_raw), "producer": pdf.metadata.get("producer"), "rawAnswers": 140, "reverseOnce": 40, "dimensions": 13, "modules": 4, "overallExact": str(overall), "overallNormExact": "705/13", "customerTextSegments": 36, "actualRingVectors": 5, "actualComparisonBars": 13, "actualNormPolyline": 1, "customerOriginalsUnchanged": original_evidence}
    (directory / "independent-oracle.json").write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({k: v for k, v in result.items() if k != "customerOriginalsUnchanged"}))


if __name__ == "__main__":
    main()