"""Independent rational/OPC/customer-text/server-PDF gate for an owned restore.

Never performs network or database writes. Input must be a fresh synthetic run.
Evidence remains explicitly SERVICE_ONLY_NOT_HTTP.
"""
import argparse
from decimal import Decimal
from fractions import Fraction
import hashlib
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET

import openpyxl
import pymupdf as fitz

# Import the unchanged, already verified oracles; never execute their old CLIs.
import importlib.util

ROOT = Path(__file__).resolve().parents[2]


def load(name, filename):
    spec = importlib.util.spec_from_file_location(name, ROOT / 'scripts/test' / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input-dir', type=Path, required=True)
    parser.add_argument('--attempt', type=int, default=1, choices=[1, 2, 3])
    args = parser.parse_args()
    directory = args.input_dir
    target = directory / ('restored-oracle.json' if args.attempt == 1 else f'restored-oracle-attempt{args.attempt}.json')
    assert not target.exists(), 'refuse to overwrite an earlier verdict'
    independent = load('mng_restored_facts', 'management-traits-four-real-pdf-oracle.py')
    pixel = load('mng_restored_pixel', 'management-traits-zero-ring-pdf-contract-test.py')
    frozen = load('mng_restored_opc', 'management-traits-zero-ring-frozen-contract-test.py')
    rows = json.loads((directory / 'reports.json').read_text('utf-8'))
    assert len(rows) == 4
    workbook_path = ROOT / 'Go-based Refactored System/configs/export-templates/management-traits-002-test-content-v1.xlsx'
    assert hashlib.sha256(workbook_path.read_bytes()).hexdigest() == 'b0498249ae057e3aa1b798922ed2d53a6ca304811943b3b41f4f64f9b024e84c'
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    original = pixel.package(pixel.TEMPLATE)
    assert pixel.receipt(pixel.TEMPLATE)['sha256'] == pixel.TEMPLATE_SHA
    envelopes = pixel.geometry(directory / '00202-candidate.pdf')
    output = {'scope': 'RESTORED_SERVICE_ONLY_NOT_HTTP', 'cases': [], 'errors': []}
    try:
        for index, row in enumerate(rows):
            name, report, run = row['name'], row['report'], row['run']
            assert row['http'] is False and row['threeSHA'] is True
            questions = {q['number']: q for q in row['questions']}
            assert sorted(questions) == list(range(1, 141))
            dims = {d['dimensionKey']: d for d in row['dimensions']}
            scores = {}
            for order, (key, module, terms, norm) in enumerate(independent.SPECS, 1):
                total = 0
                for term in terms:
                    q = questions[abs(term)]
                    final = [3, 1, 5, 1 + abs(term) % 5][index]
                    raw = 6 - final if term < 0 else final
                    options = json.loads(q['optionsSnapshot'])
                    assert q['reverse'] == (term < 0) and q['rawAnswer'] == raw and q['finalScore'] == final
                    assert next(o['raw'] for o in options if o['sourceOptionId'] == q['selectedOptionId']) == raw
                    total += final
                score = Fraction(25 * (total - len(terms)), len(terms))
                d = dims[key]
                assert (d['displayOrder'], d['scoreSum'], d['questionCount'], d['answeredCount']) == (order, total, len(terms), len(terms))
                assert Decimal(str(d['score'])) == independent.rounded(score, 6)
                assert Decimal(str(d['norm'])) == independent.rounded(norm, 6) and d['level'] == independent.level(score)
                scores[key] = score
            overall = sum(scores.values(), Fraction()) / 13
            assert run['status'] == 'completed' and run['answeredQuestionCount'] == 140
            assert Decimal(str(run['overallScore'])) == independent.rounded(overall, 6)
            assert Decimal(str(run['overallNorm'])) == independent.rounded(Fraction(705, 13), 6)
            assert run['overallLevel'] == independent.level(overall)
            for module in row['modules']:
                children = [scores[k] for k, m, _, _ in independent.SPECS if m == module['moduleKey']]
                assert module['dimensionCount'] == len(children)
                assert Decimal(str(module['score'])) == independent.rounded(sum(children, Fraction()) / len(children), 6)
            high = sorted(range(13), key=lambda n: (-scores[independent.SPECS[n][0]], n))[:3]
            low = sorted(range(13), key=lambda n: (scores[independent.SPECS[n][0]], n))[:3]
            # DataSnapshot has json:"-" in the actual model. Do not invent an
            # exported DTO/hash proof. Derive expected presentation from facts,
            # then compare each real rendered content control to customer XLSX.
            package = pixel.package(directory / (name + '.docx'))
            doc = ET.fromstring(package['word/document.xml'])
            fields = {}
            for sdt in doc.findall('.//w:sdt', pixel.NS):
                tag = sdt.find('w:sdtPr/w:tag', pixel.NS)
                assert tag is not None
                key = tag.get('{'+pixel.NS['w']+'}val')
                assert key not in fields
                fields[key] = ''.join(x.text or '' for x in sdt.findall('.//w:t', pixel.NS))
            assert len(fields) == 90
            numeric = {}
            for anchor in doc.findall('.//wp:anchor', pixel.NS):
                prop = anchor.find('wp:docPr', pixel.NS)
                if prop is not None and prop.get('title', '').endswith('.numeric-label'):
                    label = prop.get('title').removesuffix('.numeric-label')
                    texts = anchor.findall('.//w:t', pixel.NS)
                    assert len(texts) in (3, 4) and label not in numeric
                    numeric[label] = texts[1].text
            assert len(numeric) == 5
            levels = ['excellent', 'good', 'qualified', 'weak', 'insufficient']
            grade = levels.index(run['overallLevel'])
            dto = {'testOnly': True, 'overall': {'chartScore': format(independent.rounded(overall, 12), '.12f'), 'score': str(independent.rounded(overall, 2)), 'diagnosis': workbook.worksheets[3].cell(grade+2, 3).value, 'advice': workbook.worksheets[3].cell(grade+2, 4).value.split('\n')}, 'dimensions': [], 'modules': [], 'highest': [], 'lowest': []}
            assert numeric['chart.overall'] == dto['overall']['score']
            assert fields['overall.diagnosis'] == dto['overall']['diagnosis']
            for n, text in enumerate(dto['overall']['advice'], 1):
                assert fields[f'overall.advice.{n}'] == text
            for n, (key, module, _, norm) in enumerate(independent.SPECS):
                grade = levels.index(dims[key]['level'])
                d = {'key': key, 'name': dims[key]['dimensionName'], 'chartScore': format(independent.rounded(scores[key], 12), '.12f'), 'score': str(independent.rounded(scores[key], 2)), 'chartNorm': format(independent.rounded(norm, 12), '.12f'), 'norm': str(independent.rounded(norm, 2)), 'diagnosis': workbook.worksheets[1].cell(n+3, 4+grade).value, 'advice': workbook.worksheets[2].cell(n+3, 3+grade).value}
                for field in ('score', 'norm', 'diagnosis', 'advice'):
                    assert fields[f'dimension.{key}.{field}'] == d[field]
                dto['dimensions'].append(d)
            for module in row['modules']:
                children = [scores[k] for k, m, _, _ in independent.SPECS if m == module['moduleKey']]
                value = sum(children, Fraction()) / len(children)
                dto['modules'].append({'Key': module['moduleKey'], 'ChartScore': format(independent.rounded(value, 12), '.12f')})
                assert numeric['chart.module.'+module['moduleKey']] == str(independent.rounded(value, 2))
            for collection, label, ranking in [('highest', 'high', high), ('lowest', 'low', low)]:
                for slot, n in enumerate(ranking, 1):
                    key = independent.SPECS[n][0]
                    text = workbook.worksheets[0].cell(n+3, 3+levels.index(dims[key]['level'])).value
                    assert fields[f'overview.{label}.{slot}.name'] == dims[key]['dimensionName']
                    assert fields[f'overview.{label}.{slot}.text'] == text
                    dto[collection].append({'Key': key, 'Name': dims[key]['dimensionName'], 'Text': text})
            assert [x['Key'] for x in dto['highest']] == [independent.SPECS[n][0] for n in high]
            assert [x['Key'] for x in dto['lowest']] == [independent.SPECS[n][0] for n in low]
            selected = frozen.chosen_texts({'dimensions': [{'key': k, 'level': d['level']} for k, d in dims.items()], 'run': {'level': run['overallLevel']}}, dto, workbook.worksheets, independent.SPECS)
            opc = frozen.opc_contract(pixel, directory / (name + '.docx'), dto, original)
            pdf_path = directory / (name + '.pdf')
            pdf_bytes = pdf_path.read_bytes()
            assert len(pdf_bytes) == report['fileBytes'] and hashlib.sha256(pdf_bytes).hexdigest() == report['fileSha']
            pixel.SCORES[name] = frozen.scores(dto)
            translated, offsets = frozen.translated_envelopes(pixel, pdf_path, directory / '00202-candidate.pdf', envelopes)
            evidence = pixel.pdf_contract(pdf_path, name, translated, directory / (f'images-attempt{args.attempt}-' + name))
            assert not evidence['errors'], evidence['errors']
            with fitz.open(pdf_path) as pdf:
                assert len(pdf) == 9 and '24.2' in (pdf.metadata.get('producer') or '')
                text = re.sub(r'\s+', '', ''.join(p.get_text(sort=False) for p in pdf))
                assert all(re.sub(r'\s+', '', x['text']) in text for x in selected)
                assert '53.75' in text and '50.00' in text
            if index == 1:
                assert evidence['visibleCompleteGrayRings'] == 5
                assert all(r['gray']['occupiedBins'] == 360 and r['gray']['whiteHoleOccupancy'] >= .99 and r['blue']['matchingPixels'] == 0 for r in evidence['rings'])
                assert [s['text'].strip() for s in evidence['numericSlots']] == ['0.00 分'] * 5
            output['cases'].append({'name': name, 'overallExact': str(overall), 'overallDisplay': dto['overall']['score'], 'opc': opc, 'geometryOffsets': offsets, 'pdf': evidence, 'selectedRuleSegments': len(selected), 'threeSHA': True, 'independentDataSHA': 'UNVERIFIED_MODEL_JSON_OMITS_SNAPSHOT', 'presentationExpectedFromSQLComparedToActualDOCX': True, 'http': False})
            print(json.dumps({'name': name, 'exact': str(overall), 'pages': evidence['pages'], 'grayRings': evidence['visibleCompleteGrayRings'], 'texts': len(selected), 'sha256': report['fileSha']}))
    except Exception as error:
        output['errors'].append(type(error).__name__ + ': ' + str(error))
    finally:
        workbook.close()
        output['status'] = 'FAIL' if output['errors'] else 'RESTORED_PASS_NOT_HTTP'
        target.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding='utf-8')
    print('RESTORED_ORACLE=' + output['status'] + ' ERRORS=' + repr(output['errors']))
    return 1 if output['errors'] else 0


if __name__ == '__main__':
    raise SystemExit(main())