"""Offline independent oracle for this release's real UI and native-saveAs files.

No network, DB writes, renderer execution, old receipt edits, or API-byte downloads.
"""
import argparse
from decimal import Decimal
from fractions import Fraction
import hashlib
import importlib.util
import json
from pathlib import Path

import openpyxl
import pymupdf as fitz

ROOT = Path(__file__).resolve().parents[2]


def load(name, filename):
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts/test" / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("directory", type=Path)
    args = parser.parse_args()
    directory = args.directory.resolve()
    assert directory.parent == (ROOT / "scripts/test/results").resolve()
    assert directory.name.startswith("uf054-ui-")
    output = directory / "native-independent-oracle.json"
    assert not output.exists(), "Prior receipts must never be overwritten"
    summary = json.loads((directory / "summary.json").read_text("utf-8"))
    assert summary["stages"]["native"] == "completed"
    assert summary["stages"]["cleanup"] == "completed"
    independent = load("uf054_rational", "management-traits-four-real-pdf-oracle.py")
    pixel = load("uf054_pixels", "management-traits-zero-ring-pdf-contract-test.py")
    metadata = json.loads((directory / "native-db-independent-match.json").read_text("utf-8"))["metadata"]
    workbook_path = ROOT / "Go-based Refactored System/configs/export-templates/management-traits-002-test-content-v1.xlsx"
    assert hashlib.sha256(workbook_path.read_bytes()).hexdigest() == "b0498249ae057e3aa1b798922ed2d53a6ca304811943b3b41f4f64f9b024e84c"
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    levels = ["excellent", "good", "qualified", "weak", "insufficient"]
    receipts = []
    envelopes = pixel.geometry(directory / "00202-candidate.pdf")
    for label in ["00201-candidate", "00201-tester", "00202-candidate", "00202-tester"]:
        case = next(x for x in summary["cases"] if x["label"] == label)
        native = next(x for x in summary["native"] if x["case"] == label)
        meta = next(x for x in metadata if x["id"] == native["reportId"])
        assert native["downloadEvent"] and native["saveAsCalls"] == 1
        assert hashlib.sha256(meta["dataRaw"].encode("utf-8")).hexdigest() == meta["dataSHA"]
        dto = json.loads(meta["dataRaw"])
        assert dto["runId"] == native["runId"] and dto["testOnly"]
        mapping = json.loads((directory / ("mapping-" + label + "-1.json")).read_text("utf-8"))
        questions = [json.loads(x) for x in mapping["stdout"].splitlines() if x.strip()]
        by_v = {x["v"]: x for x in questions}
        answers = {x["questionId"]: x for x in case["answers"]}
        assert sorted(by_v) == list(range(1, 141)) and len(answers) == 140
        assert all(x["http"] == 200 and x["code"] == 0 for x in answers.values())
        detail = json.loads((directory / ("detail-" + label + ".json")).read_text("utf-8"))
        dims = {x["Key"]: x for x in detail["Result"]["Dimensions"]}
        scores = {}
        for key, module, terms, norm in independent.SPECS:
            total = 0
            for term in terms:
                q = by_v[abs(term)]
                assert bool(q["reverse"]) == (term < 0)
                answer = answers[q["id"]]
                raw = answer["raw"]
                expected = {"fifty": 3, "zero": 5 if term < 0 else 1,
                            "hundred": 1 if term < 0 else 5,
                            "mixed": 1 + abs(term) % 5}[case["pattern"]]
                assert raw == expected
                option = next(x for x in q["options"] if x["sourceOptionId"] == answer["optionId"])
                assert option["raw"] == raw
                total += 6 - raw if term < 0 else raw
            score = Fraction(25 * (total - len(terms)), len(terms))
            scores[key] = score
            d = dims[key]
            assert (d["ScoreSum"], d["QuestionCount"], d["AnsweredCount"]) == (total, len(terms), len(terms))
            assert Fraction(d["Score"]) == score and Fraction(d["Norm"]) == norm
            assert d["Level"] == independent.level(score)
        overall = sum(scores.values(), Fraction()) / 13
        assert Fraction(detail["Result"]["OverallScore"]) == overall
        assert Fraction(detail["Result"]["OverallNorm"]) == Fraction(705, 13)
        assert dto["overall"]["score"] == str(independent.rounded(overall, 2))
        for module in detail["Result"]["Modules"]:
            children = [scores[key] for key, group, _, _ in independent.SPECS if group == module["Key"]]
            assert module["DimensionCount"] == len(children)
            assert Fraction(module["Score"]) == sum(children, Fraction()) / len(children)
        selected = []
        for index, d in enumerate(dto["dimensions"]):
            key = independent.SPECS[index][0]
            grade = levels.index(dims[key]["Level"])
            assert d["key"] == key
            assert d["score"] == str(independent.rounded(scores[key], 2))
            assert d["diagnosis"] == workbook.worksheets[1].cell(index + 3, 4 + grade).value
            assert d["advice"] == workbook.worksheets[2].cell(index + 3, 3 + grade).value
            selected += [d["diagnosis"], d["advice"]]
        high = sorted(range(13), key=lambda n: (-scores[independent.SPECS[n][0]], n))[:3]
        low = sorted(range(13), key=lambda n: (scores[independent.SPECS[n][0]], n))[:3]
        for name, ranking in [("highest", high), ("lowest", low)]:
            assert [x["Key"] for x in dto[name]] == [independent.SPECS[n][0] for n in ranking]
            for x, index in zip(dto[name], ranking, strict=True):
                grade = levels.index(dims[x["Key"]]["Level"])
                assert x["Text"] == workbook.worksheets[0].cell(index + 3, 3 + grade).value
                selected.append(x["Text"])
        grade = levels.index(independent.level(overall))
        assert dto["overall"]["diagnosis"] == workbook.worksheets[3].cell(grade + 2, 3).value
        assert dto["overall"]["advice"] == workbook.worksheets[3].cell(grade + 2, 4).value.split("\n")
        selected += [dto["overall"]["diagnosis"], *dto["overall"]["advice"]]
        data = (directory / (label + ".pdf")).read_bytes()
        assert len(data) == native["bytes"] == meta["bytes"]
        assert hashlib.sha256(data).hexdigest() == native["sha256"] == meta["sha"]
        assert data.startswith(b"%PDF-")
        with fitz.open(stream=data, filetype="pdf") as pdf:
            assert len(pdf) == 9 and all(p.get_text().strip() for p in pdf)
            assert all(abs(p.rect.width - 595.3) < 2 and abs(p.rect.height - 841.9) < 2 for p in pdf)
            text = independent.compact("".join(p.get_text(sort=False) for p in pdf))
            assert "TEST" in text and "不可作为人才决策依据" in text and "\ufffd" not in text
            assert all(independent.compact(x) in text for x in selected)
            images = directory / ("native-pages-" + label)
            images.mkdir(exist_ok=False)
            contact = fitz.open()
            sheet = contact.new_page(width=600, height=849)
            for i, page in enumerate(pdf):
                pix = page.get_pixmap(matrix=fitz.Matrix(1, 1))
                pix.save(images / ("page-" + str(i + 1) + ".png"))
                x, y = i % 3 * 200, i // 3 * 283
                sheet.insert_image(fitz.Rect(x, y, x + 200, y + 283), stream=pix.tobytes("png"))
            sheet.get_pixmap().save(images / "all-pages.png")
            zero_rings = []
            if case["pattern"] == "zero":
                for key, rectangle in envelopes.items():
                    ring = pixel.annulus(pdf[2], rectangle, pixel.GRAY)
                    blue = pixel.annulus(pdf[2], rectangle, pixel.BLUE)
                    assert ring["occupiedBins"] == 360 and ring["whiteHoleOccupancy"] >= .99
                    assert blue["matchingPixels"] == 0
                    zero_rings.append({"key": key, "occupiedBins": 360, "whiteHole": ring["whiteHoleOccupancy"], "bluePixels": 0})
        receipts.append({"case": label, "overallExact": str(overall), "overallDisplay": str(independent.rounded(overall, 2)),
                         "dimensions": 13, "modules": 4, "answers": 140, "reverseOnce": 40,
                         "nativeDBSHA": native["sha256"], "pages": 9, "selectedCustomerSegments": len(selected),
                         "zeroRings": zero_rings, "internalDOCX": "NOT_CAPTURED", "mainRace": "UNPROVEN"})
    output.write_text(json.dumps({"scope": "NEW_REAL_UI_NATIVE_OFFLINE_INDEPENDENT", "exit": 0, "cases": receipts}, indent=2), "utf-8")
    print(json.dumps({"exit": 0, "nativeReports": 4, "answerFacts": 560, "pages": 36,
                      "customerSegments": sum(x["selectedCustomerSegments"] for x in receipts), "mainRace": "UNPROVEN"}))


if __name__ == "__main__":
    main()