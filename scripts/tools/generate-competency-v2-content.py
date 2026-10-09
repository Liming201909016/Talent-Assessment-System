#!/usr/bin/env python3
import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

PRODUCT = 'competency-frontline-phase1-v2'
SCORING = 'competency-phase1-scoring-v2'
CONTENT = 'competency-phase1-content-v2'
TEMPLATE = 'competency-phase1-report-v2'
AUDIENCE = 'frontline_employee'
DISCLAIMER = '本报告基于受测者在本次胜任力测评中的作答结果生成，用于辅助了解其当前胜任力表现及发展方向。测评结果可能受到作答状态、岗位经历和测评环境等因素影响，仅供人才发展、培训与管理决策参考，不应作为招聘、晋升、淘汰或其他重大人事决定的唯一依据。使用者应结合岗位要求、工作绩效、结构化访谈及其他有效信息进行综合判断。未经授权，不得复制、传播或向无关人员披露本报告内容。'
VALIDITY = {
    'good': '本次测评作答效度良好，结果具有较好的参考价值。测评结果仍应结合实际工作表现、行为观察、访谈及其他评价信息进行综合解读，不宜作为人才决策的唯一依据。',
    'questionable': '该受测者存在掩饰真实想法的可能性，本次测评结果请谨慎解读，必要时可重新开展评估。',
}
DIMENSIONS = [
    ('A1-01', 'competency-logical-reasoning'), ('A1-02', 'competency-plan-execution'),
    ('A1-03', 'competency-digital-application'), ('A1-04', 'competency-achievement-orientation'),
    ('A1-05', 'competency-continuous-learning'), ('B1-01', 'competency-communication'),
    ('B1-02', 'competency-cooperation'), ('C1-01', 'competency-truth-pragmatism'),
    ('C1-02', 'competency-self-discipline'), ('C1-03', 'competency-dedication'),
]
LEVELS = ['not_qualified', 'weak', 'qualified', 'good', 'excellent']
MODULES = ['task_management', 'interpersonal_management', 'self_management']
COMPARISONS = {'standout': '优势突出', 'above_norm': '略高于常模分', 'at_norm': '与常模分持平', 'below_norm': '低于常模分'}


def clean(value):
    return ' '.join(str(value).replace('\u3000', ' ').split())


def add(rows, content_type, identity, condition, content, source):
    text = clean(content)
    if not text:
        raise ValueError(f'blank content: {content_type}/{identity}/{condition}')
    rows.append({'contentType': content_type, 'identity': identity, 'condition': condition, 'content': text, 'source': source})


def build(workbook):
    wb = load_workbook(workbook, data_only=True)
    rows = []
    dimension_sheet = wb['等级评价']
    by_code = {}
    for values in dimension_sheet.iter_rows(values_only=True):
        code_name = values[2] if len(values) > 2 else None
        if isinstance(code_name, str) and '-' in code_name:
            by_code[clean(code_name).split(' ')[0]] = values
    for code, dimension_id in DIMENSIONS:
        values = by_code.get(code)
        if values is None:
            raise ValueError(f'dimension row missing: {code}')
        for index, level in enumerate(LEVELS):
            full = values[5 + index * 2]
            short = values[6 + index * 2]
            add(rows, 'dimension', dimension_id, level, full, f'等级评价:{code}:L{index + 1}:full')
            selected_type = 'strength' if level in ('good', 'excellent') else 'development'
            add(rows, selected_type, dimension_id, level, short, f'等级评价:{code}:L{index + 1}:short')

    overall_sheet = wb['总体评价']
    overall_rows = [row for row in overall_sheet.iter_rows(min_row=3, max_row=7, values_only=True)]
    if len(overall_rows) != 5:
        raise ValueError('overall rows are incomplete')
    # Workbook order is excellent -> not qualified; normalize to internal level order.
    for level, values in zip(reversed(LEVELS), overall_rows):
        add(rows, 'overall', '', level, values[2], f'总体评价:{values[0]}:assessment')
        add(rows, 'overall_advice', '', level, values[3], f'总体评价:{values[0]}:advice')

    for module in MODULES:
        for condition, text in COMPARISONS.items():
            add(rows, 'module_comparison', module, condition, text, f'三类模块评价:{module}:{condition}')
    for condition, text in VALIDITY.items():
        add(rows, 'validity', '', condition, text, f'competency-phase1-content-v1-approved:{condition}')

    rows.sort(key=lambda row: (row['contentType'], row['identity'], row['condition']))
    keys = {(row['contentType'], row['identity'], row['condition']) for row in rows}
    if len(rows) != 124 or len(keys) != 124:
        raise ValueError(f'v2 content cardinality mismatch: rows={len(rows)}, keys={len(keys)}')
    return rows


def sql_hex(value):
    return "CONVERT(UNHEX('%s') USING utf8mb4)" % value.encode('utf-8').hex()


def render_sql(rows, question_sha, content_sha):
    output = [
        'SET NAMES utf8mb4;', 'BEGIN;',
        "DELETE FROM `el_competency_report_text` WHERE `content_version`='competency-phase1-content-v2' AND `audience`='frontline_employee';",
        "DELETE FROM `el_competency_report_content_package` WHERE `product_version`='competency-frontline-phase1-v2' AND `scoring_version`='competency-phase1-scoring-v2' AND `content_version`='competency-phase1-content-v2' AND `template_version`='competency-phase1-report-v2' AND `audience`='frontline_employee';",
    ]
    for index, row in enumerate(rows, 1):
        identity = row['identity']
        row_id = f'phase1-v2-{index:03d}'
        output.append(
            'INSERT INTO `el_competency_report_text` '
            '(`id`,`content_version`,`audience`,`content_type`,`dimension_id`,`level_code`,`content`,`disclaimer`,`is_temporary`,`status`,`create_time`,`update_time`) VALUES '
            f"('{row_id}','{CONTENT}','{AUDIENCE}','{row['contentType']}','{identity}','{row['condition']}',"
            f"{sql_hex(row['content'])},{sql_hex(DISCLAIMER)},0,1,NOW(),NOW());"
        )
    output.append(
        'INSERT INTO `el_competency_report_content_package` '
        '(`id`,`product_version`,`scoring_version`,`content_version`,`template_version`,`audience`,`approval_status`,`content_approved_by`,`content_approved_at`,`psychometric_approved_by`,`psychometric_approved_at`,`question_source_sha256`,`content_source_sha256`,`effective_environment`,`disclaimer`,`create_time`,`update_time`) VALUES '
        f"('phase1-v2-content-draft','{PRODUCT}','{SCORING}','{CONTENT}','{TEMPLATE}','{AUDIENCE}','draft','',NULL,'',NULL,'{question_sha}','{content_sha}','','{DISCLAIMER}',NOW(),NOW());"
    )
    output.extend([
        "SELECT COUNT(*) INTO @v2_text_count FROM `el_competency_report_text` WHERE `content_version`='competency-phase1-content-v2' AND `audience`='frontline_employee';",
        "SET @sql=IF(@v2_text_count=124,'SELECT ''v2 draft content cardinality valid''','SELECT * FROM `__v2_draft_content_cardinality_invalid__`');",
        'PREPARE stmt FROM @sql; EXECUTE stmt; DEALLOCATE PREPARE stmt;', 'COMMIT;', '',
    ])
    return '\n'.join(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--json', type=Path, required=True)
    parser.add_argument('--sql', type=Path, required=True)
    args = parser.parse_args()
    source = args.workbook.read_bytes()
    rows = build(args.workbook)
    canonical = json.dumps({'version': CONTENT, 'audience': AUDIENCE, 'disclaimer': DISCLAIMER, 'rows': rows}, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode()
    content_sha = hashlib.sha256(canonical).hexdigest()
    question_sha = hashlib.sha256(source).hexdigest()
    payload = json.loads(canonical.decode())
    payload['questionSourceSHA256'] = question_sha
    payload['contentSourceSHA256'] = content_sha
    args.json.parent.mkdir(parents=True, exist_ok=True)
    args.sql.parent.mkdir(parents=True, exist_ok=True)
    args.json.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    args.sql.write_text(render_sql(rows, question_sha, content_sha), encoding='utf-8')
    print(f'V2_CONTENT_DRAFT_PASS rows={len(rows)} questionSHA={question_sha} contentSHA={content_sha}')


if __name__ == '__main__':
    main()
